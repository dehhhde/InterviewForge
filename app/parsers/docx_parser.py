"""Word 解析 python-docx：段落 + 表格按文档流拉平〔阶段 5 · §8.3〕

为什么必须按 `body` 的子元素顺序遍历：`document.paragraphs` 与 `document.tables` 是两个
独立列表，先段落再表格会把"简历自评表"整块挪到全文末尾，脱离它上面的段落标题；
实测本仓库的架构文档里第 10 个子元素就是表格。

只取 OOXML（`.docx`）。97-2003 的 `.doc` 是二进制复合文档，python-docx 读不了，
由 registry 按扩展名直接拒为"不支持的类型"，不在这里伪装成功。
"""

import re
import zipfile
from collections.abc import Iterator
from pathlib import Path

from docx.api import Document as open_document
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.exceptions.base import ParseError
from app.parsers.base import BaseParser, ParsedDocument

#: 同一行单元格的分隔符。用制表符而不是 ` | `：Word 里复制表格到纯文本就是制表符分隔，
#: 而正文与代码块里出现 `|` 很常见（`dict | None`），用竖线会让下游误判那是表格行。
CELL_SEPARATOR = "	"
_PAGES_RE = re.compile(r"<Pages>(\d+)</Pages>")


class DocxParser(BaseParser):
    extensions = (".docx",)

    def parse(self, path: Path) -> ParsedDocument:
        try:
            document = open_document(str(path))
        except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile) as exc:
            raise ParseError(
                "DOCX 打不开：可能是 .doc 老格式、已损坏或被加密",
                code="parse_bad_docx",
                details={"path": str(path), "reason": str(exc)},
            ) from exc

        blocks = _header_lines(document) + _body_lines(document)
        raw_text = "\n".join(blocks)
        if not raw_text.strip():
            raise ParseError(
                "DOCX 没有可提取的文字", code="parse_empty_docx", details={"path": str(path)}
            )

        return ParsedDocument(
            raw_text=raw_text,
            page_count=_declared_page_count(path),
            is_scanned=False,
        )


def _header_lines(document: DocxDocument) -> list[str]:
    """页眉按视觉顺序放在正文前——中文简历习惯把姓名与联系方式写在页眉里。

    页脚不取：那里通常是页码与版权声明，抽进正文只会喂噪声。
    """
    lines: list[str] = []
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            text = paragraph.text.strip()
            if text and text not in lines:
                lines.append(text)
    return lines


def _body_lines(document: DocxDocument) -> list[str]:
    lines: list[str] = []
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            text = Paragraph(element, document).text.strip()
            if text:
                lines.append(text)
        elif element.tag == qn("w:tbl"):
            lines.extend(_table_lines(Table(element, document)))
    return lines


def _table_lines(table: Table) -> Iterator[str]:
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        line = CELL_SEPARATOR.join(cell for cell in cells if cell)
        if line:
            yield line


def _declared_page_count(path: Path) -> int:
    """docx 的分页由 Word 渲染时才决定，但 Word 另存时会把上次渲染的页数写进
    `docProps/app.xml`。读得到就用，读不到返回 1。

    这个数只能看不能信：实测本仓三份设计文档的 `app.xml` 都写着 1 页，实际最长的一份
    接近 700 行。因此它只进日志与 trace，任何质量阈值都不许依赖它。
    """
    try:
        with zipfile.ZipFile(path) as archive:
            app = archive.read("docProps/app.xml").decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return 1
    match = _PAGES_RE.search(app)
    return int(match.group(1)) if match else 1
